"""DOCX lesson-plan generator based on python-docx."""

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


def generate_docx(lesson_ir: dict) -> bytes:
    """Render a lesson outline/IR into a valid downloadable DOCX."""
    document = Document()
    title = lesson_ir.get("title") or lesson_ir.get("topic") or "未命名教案"
    heading = document.add_heading(str(title), level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subject = lesson_ir.get("subject", "")
    grade = lesson_ir.get("grade", "")
    if subject or grade:
        document.add_paragraph(f"科目：{subject}    年级：{grade}")

    slots = lesson_ir.get("slots") or lesson_ir
    for label, key in (("学习目标", "objectives"), ("教学重点", "key_points"), ("先备知识", "prerequisites")):
        values = slots.get(key) or []
        if values:
            document.add_heading(label, level=1)
            for value in values:
                document.add_paragraph(str(value), style="List Bullet")

    sections = lesson_ir.get("sections") or []
    if sections:
        document.add_heading("课程大纲", level=1)
        for index, section in enumerate(sections, start=1):
            document.add_heading(f"{index}. {section.get('title', '未命名章节')}", level=2)
            duration = section.get("duration_minutes", 0)
            slide_count = section.get("slide_count", 0)
            document.add_paragraph(f"时长：{duration} 分钟    建议页数：{slide_count} 页")
            for bullet in section.get("bullets", []):
                document.add_paragraph(str(bullet), style="List Bullet")

    output = BytesIO()
    document.save(output)
    return output.getvalue()
