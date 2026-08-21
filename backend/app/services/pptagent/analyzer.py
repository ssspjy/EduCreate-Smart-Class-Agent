"""PPTX 参考课件结构分析。"""

from pathlib import Path
from fastapi import HTTPException, status
from pptx import Presentation


def analyze_reference(path: str, material_id: str = "", filename: str = "") -> dict:
    file_path = Path(path).resolve()
    if not file_path.is_file():
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="参考课件文件不存在")
    try:
        prs = Presentation(str(file_path))
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail=f"PPTX 无法读取：{exc}") from exc
    slides = []
    lengths = []
    for index, slide in enumerate(prs.slides, start=1):
        texts = [shape.text.strip() for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()]
        chars = sum(len(text) for text in texts)
        lengths.append(chars)
        slides.append({"index": index, "text_blocks": len(texts), "characters": chars, "has_title": bool(texts and len(texts[0]) <= 120)})
    average = round(sum(lengths) / len(lengths), 1) if lengths else 0
    density = "sparse" if average < 80 else "balanced" if average < 220 else "dense"
    return {"material_id": material_id, "filename": filename or file_path.name, "slide_count": len(prs.slides), "average_characters": average, "density": density, "recommended_style": "minimal" if density == "sparse" else "classic" if density == "balanced" else "modern", "slides": slides, "warnings": [] if slides else ["参考课件没有可分析的幻灯片"]}


def analyze_reference_ppt(ppt_path: str) -> dict:
    return analyze_reference(ppt_path)
