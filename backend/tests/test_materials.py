"""Material upload and persistence tests."""

import asyncio
from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from docx import Document
from pypdf import PdfWriter

from app.db import SessionLocal
from app.models import Chunk, Material
from app.core.config import get_settings
from app.services import material_service
from app.services.parsers import parser as parser_module
from app.services.parsers.parser import ParsedChunk, ParseResult
from app.services.parsers.ocr import OCRResult
from app.services.parsers.video import TranscriptSegment, VideoParseResult


def test_text_splitter_removes_nul_characters() -> None:
    chunks = parser_module._split_text("第一段\x00仍是同一段\n第二段")

    assert chunks == ["第一段仍是同一段\n第二段"]


def test_failed_upload_stream_removes_partial_file(tmp_path: Path) -> None:
    class BrokenUpload:
        async def read(self, _chunk_size: int) -> bytes:
            if not hasattr(self, "called"):
                self.called = True
                return b"partial"
            raise RuntimeError("stream read failed")

        async def seek(self, _position: int) -> None:
            return None

    target = tmp_path / "uploads" / "partial.bin"
    try:
        asyncio.run(material_service._save_upload(BrokenUpload(), target, max_bytes=1024))
    except RuntimeError as exc:
        assert str(exc) == "stream read failed"
    else:
        raise AssertionError("broken upload should raise")

    assert not target.exists()


def test_upload_image_is_persisted_without_fake_chunks(client: TestClient) -> None:
    response = client.post(
        "/api/v1/materials/upload",
        files={"file": ("diagram.png", b"not-a-real-image-but-stored", "image/png")},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["filename"] == "diagram.png"
    assert payload["status"] == "uploaded"
    assert payload["chunk_count"] == 0
    assert "OCR" in payload["error_message"]

    detail = client.get(f"/api/v1/materials/{payload['file_id']}")
    assert detail.status_code == 200
    assert detail.json()["chunks"] == []


def test_upload_image_uses_ocr_when_text_is_recognized(
    client: TestClient,
    monkeypatch,
) -> None:
    def fake_ocr_image(*args, **kwargs) -> OCRResult:
        return OCRResult(text_by_page={1: "图片识别：阿基米德原理"})

    monkeypatch.setattr(parser_module, "ocr_image", fake_ocr_image)
    response = client.post(
        "/api/v1/materials/upload",
        files={"file": ("formula.png", b"mock-image", "image/png")},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "parsed"
    chunks = client.get(f"/api/v1/materials/{payload['file_id']}/chunks").json()
    assert chunks[0]["content"] == "图片识别：阿基米德原理"
    assert chunks[0]["modality"] == "ocr"


def test_upload_docx_extracts_text_chunks(client: TestClient) -> None:
    document = Document()
    document.add_paragraph("光的折射发生在两种介质的交界面。")
    document.add_paragraph("入射光线、折射光线和法线在同一平面内。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "折射率"
    table.cell(0, 1).text = "n"
    table.cell(1, 0).text = "公式"
    table.cell(1, 1).text = "sin i / sin r"
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)

    response = client.post(
        "/api/v1/materials/upload",
        files={
            "file": (
                "refraction.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "parsed"
    assert payload["chunk_count"] >= 1

    with SessionLocal() as db:
        stored_chunk = db.query(Chunk).filter(Chunk.material_id == payload["file_id"]).first()
        assert stored_chunk is not None
        assert stored_chunk.embedding is not None
        assert len(stored_chunk.embedding) == 1024

    chunks = client.get(f"/api/v1/materials/{payload['file_id']}/chunks")
    assert chunks.status_code == 200
    joined = "\n".join(chunk["content"] for chunk in chunks.json())
    assert "光的折射" in joined
    assert "同一平面" in joined
    assert "折射率" in joined


def test_upload_markdown_extracts_text_chunks(client: TestClient) -> None:
    response = client.post(
        "/api/v1/materials/upload",
        files={"file": ("notes.md", "## 浮力\n阿基米德原理是重点。", "text/markdown")},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "parsed"
    assert payload["chunk_count"] == 1

    chunks = client.get(f"/api/v1/materials/{payload['file_id']}/chunks").json()
    assert "阿基米德原理" in chunks[0]["content"]

    search = client.post(
        "/api/v1/knowledge/search",
        json={"query": "阿基米德原理", "top_k": 3, "material_ids": [payload["file_id"]]},
    )
    assert search.status_code == 200
    assert search.json()[0]["source"] == "notes.md"


def test_delete_material_removes_record_and_file(client: TestClient) -> None:
    response = client.post(
        "/api/v1/materials/upload",
        files={"file": ("notes.txt", "可删除材料", "text/plain")},
    )
    assert response.status_code == 200
    material_id = response.json()["file_id"]

    deleted = client.delete(f"/api/v1/materials/{material_id}")
    assert deleted.status_code == 204
    assert client.get(f"/api/v1/materials/{material_id}").status_code == 404


def test_blank_pdf_reports_ocr_warning(client: TestClient) -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    buffer = BytesIO()
    writer.write(buffer)

    response = client.post(
        "/api/v1/materials/upload",
        files={"file": ("scan.pdf", buffer.getvalue(), "application/pdf")},
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "uploaded"
    assert "OCR" in payload["error_message"]


def test_scanned_pdf_uses_ocr_and_persists_page_metadata(
    client: TestClient,
    monkeypatch,
) -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    buffer = BytesIO()
    writer.write(buffer)

    def fake_ocr_pdf_pages(*args, **kwargs) -> OCRResult:
        return OCRResult(text_by_page={1: "扫描识别：光的折射定律"})

    monkeypatch.setattr(parser_module, "ocr_pdf_pages", fake_ocr_pdf_pages)
    response = client.post(
        "/api/v1/materials/upload",
        files={"file": ("scan.pdf", buffer.getvalue(), "application/pdf")},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "parsed"
    assert payload["chunk_count"] == 1
    assert payload["error_message"] is None

    chunks = client.get(f"/api/v1/materials/{payload['file_id']}/chunks").json()
    assert chunks[0]["content"] == "扫描识别：光的折射定律"
    assert chunks[0]["page_ref"] == 1
    assert chunks[0]["modality"] == "ocr"


def test_export_docx_returns_valid_word_document(client: TestClient) -> None:
    response = client.post(
        "/api/v1/exports/docx",
        json={
            "title": "光的折射教案",
            "subject": "物理",
            "grade": "初中",
            "sections": [
                {"title": "折射率", "bullets": ["理解折射率公式"], "duration_minutes": 10, "slide_count": 1}
            ],
        },
    )
    assert response.status_code == 200
    download = client.get(response.json()["url"])
    assert download.status_code == 200
    parsed = Document(BytesIO(download.content))
    assert "光的折射教案" in "\n".join(p.text for p in parsed.paragraphs)


def test_upload_filename_path_is_sanitized(client: TestClient) -> None:
    response = client.post(
        "/api/v1/materials/upload",
        files={"file": ("..\\..\\escape.png", b"image-placeholder", "image/png")},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["filename"] == "escape.png"


def test_upload_video_persists_timestamped_transcript_chunks(
    client: TestClient,
    monkeypatch,
) -> None:
    def fake_parse_video(*args, **kwargs) -> VideoParseResult:
        return VideoParseResult(
            duration_seconds=4.2,
            segments=[TranscriptSegment(start=0.4, end=2.1, text="光的折射定律")],
        )

    monkeypatch.setattr(parser_module, "parse_video", fake_parse_video)
    response = client.post(
        "/api/v1/materials/upload",
        files={"file": ("lesson.mp4", b"mock-video", "video/mp4")},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "parsed"
    chunks = client.get(f"/api/v1/materials/{payload['file_id']}/chunks").json()
    assert chunks[0]["content"] == "[0.4s–2.1s] 光的折射定律"
    assert chunks[0]["media_ref"].endswith(".mp4#t=0.4-2.1")
    assert payload["file_id"] in chunks[0]["media_ref"]
    assert chunks[0]["modality"] == "transcript"


def test_upload_audio_persists_transcript_chunks(client: TestClient, monkeypatch) -> None:
    def fake_parse_audio(*args, **kwargs) -> VideoParseResult:
        return VideoParseResult(
            duration_seconds=2.0,
            segments=[TranscriptSegment(start=0.0, end=1.5, text="语音输入测试")],
        )

    monkeypatch.setattr(parser_module, "parse_audio", fake_parse_audio)
    response = client.post(
        "/api/v1/materials/upload",
        files={"file": ("voice.webm", b"mock-audio", "audio/webm")},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "parsed"
    chunks = client.get(f"/api/v1/materials/{payload['file_id']}/chunks").json()
    assert chunks[0]["content"] == "[0.0s–1.5s] 语音输入测试"
    assert chunks[0]["modality"] == "transcript"


def test_corrupt_supported_file_is_recorded_as_failed(client: TestClient) -> None:
    response = client.post(
        "/api/v1/materials/upload",
        files={"file": ("broken.pdf", b"not a pdf", "application/pdf")},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "failed"
    assert payload["chunk_count"] == 0
    assert payload["error_message"]

    detail = client.get(f"/api/v1/materials/{payload['file_id']}")
    assert detail.status_code == 200
    assert detail.json()["status"] == "failed"


def test_rejects_unsupported_upload_extension(client: TestClient) -> None:
    response = client.post(
        "/api/v1/materials/upload",
        files={"file": ("unsafe.exe", b"binary", "application/octet-stream")},
    )

    assert response.status_code == 400
    assert "unsupported file extension" in response.json()["detail"]


def test_generate_lesson_persists_workspace(client: TestClient) -> None:
    response = client.post(
        "/api/v1/lessons",
        json={"title": "教案-测试", "subject": "物理", "grade": "初中三年级", "topic": "浮力"},
    )

    assert response.status_code == 200
    lesson = response.json()
    assert lesson["id"]
    assert lesson["status"] == "draft"

    lessons = client.get("/api/v1/lessons")
    assert lessons.status_code == 200
    assert any(item["id"] == lesson["id"] for item in lessons.json())


def test_upload_file_is_stored_inside_test_runtime(client: TestClient) -> None:
    response = client.post(
        "/api/v1/materials/upload",
        files={"file": ("diagram.png", b"stored", "image/png")},
    )

    assert response.status_code == 200
    payload = response.json()
    stored_files = list((Path(__file__).resolve().parent / "_tmp" / "uploads").glob("**/*.png"))
    assert len(stored_files) == 1
    assert payload["file_id"] in str(stored_files[0])


def test_async_upload_is_processed_by_shared_worker_path(
    client: TestClient,
    monkeypatch,
) -> None:
    queued: list[tuple[str, str | None]] = []
    monkeypatch.setattr(get_settings(), "material_async_enabled", True)
    monkeypatch.setattr(
        material_service,
        "enqueue_material_parse",
        lambda material_id, task_id: queued.append((material_id, task_id)),
    )

    response = client.post(
        "/api/v1/materials/upload",
        files={"file": ("queued.txt", "后台解析材料", "text/plain")},
    )
    payload = response.json()
    assert response.status_code == 200
    assert payload["status"] == "queued"
    assert payload["parse_progress"] == 0
    assert payload["can_cancel"] is True
    assert queued[0][0] == payload["file_id"]

    with SessionLocal() as db:
        parsed = asyncio.run(material_service.parse_material_record(db, payload["file_id"]))
    assert parsed.status == "parsed"
    assert parsed.parse_progress == 100
    assert parsed.chunk_count == 1


def test_queued_material_can_be_cancelled_and_deleted(
    client: TestClient,
    monkeypatch,
) -> None:
    revoked: list[str] = []
    monkeypatch.setattr(get_settings(), "material_async_enabled", True)
    monkeypatch.setattr(material_service, "enqueue_material_parse", lambda *_: None)
    monkeypatch.setattr(material_service, "revoke_material_parse", revoked.append)

    uploaded = client.post(
        "/api/v1/materials/upload",
        files={"file": ("cancel.txt", "待取消", "text/plain")},
    ).json()
    blocked_delete = client.delete(f"/api/v1/materials/{uploaded['file_id']}")
    assert blocked_delete.status_code == 409

    cancelled = client.post(f"/api/v1/materials/{uploaded['file_id']}/cancel")
    assert cancelled.status_code == 200
    assert cancelled.json()["status"] == "cancelled"
    assert cancelled.json()["can_cancel"] is False
    assert len(revoked) == 1

    assert client.delete(f"/api/v1/materials/{uploaded['file_id']}").status_code == 204


def test_worker_observes_cancellation_committed_during_parse(
    client: TestClient,
    monkeypatch,
) -> None:
    monkeypatch.setattr(get_settings(), "material_async_enabled", True)
    monkeypatch.setattr(material_service, "enqueue_material_parse", lambda *_: None)
    uploaded = client.post(
        "/api/v1/materials/upload",
        files={"file": ("slow.txt", "慢速材料", "text/plain")},
    ).json()

    async def fake_parse(*_args, **_kwargs) -> ParseResult:
        with SessionLocal() as other_db:
            material = other_db.get(Material, uploaded["file_id"])
            assert material is not None
            material.cancel_requested = True
            material.status = "cancelled"
            other_db.commit()
        return ParseResult(chunks=[ParsedChunk(content="不应持久化")], warnings=[])

    monkeypatch.setattr(material_service, "parse", fake_parse)
    with SessionLocal() as db:
        result = asyncio.run(material_service.parse_material_record(db, uploaded["file_id"]))

    assert result.status == "cancelled"
    assert result.chunk_count == 0


def test_running_material_stays_undeletable_until_worker_confirms_cancel(
    client: TestClient,
    monkeypatch,
) -> None:
    monkeypatch.setattr(get_settings(), "material_async_enabled", True)
    monkeypatch.setattr(material_service, "enqueue_material_parse", lambda *_: None)
    monkeypatch.setattr(material_service, "revoke_material_parse", lambda *_: None)
    uploaded = client.post(
        "/api/v1/materials/upload",
        files={"file": ("running.txt", "运行中材料", "text/plain")},
    ).json()
    with SessionLocal() as db:
        material = db.get(Material, uploaded["file_id"])
        assert material is not None
        material.status = "parsing"
        db.commit()

    cancelling = client.post(f"/api/v1/materials/{uploaded['file_id']}/cancel")
    assert cancelling.status_code == 200
    assert cancelling.json()["status"] == "cancelling"
    assert client.delete(f"/api/v1/materials/{uploaded['file_id']}").status_code == 409

    with SessionLocal() as db:
        confirmed = asyncio.run(material_service.parse_material_record(db, uploaded["file_id"]))
    assert confirmed.status == "cancelled"
    assert client.delete(f"/api/v1/materials/{uploaded['file_id']}").status_code == 204
